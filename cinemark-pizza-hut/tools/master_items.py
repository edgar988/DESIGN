#!/usr/bin/env python3
"""Pizza Hut master item numbers: the same number for a product on every PH project (Edgar).

config/master_items.json holds the master list (no prices: those stay in the AutoQuotes export on
Drive). Each item says how a drawing block is recognised as it:
  match   regexes tried against the block's AutoQuotes model, KCL hyperlink model and block name
          (the model itself is always tried too)
  size    [width, depth] inches, for stand-in blocks matched by footprint (fabricated worktables drawn
          with Advance Tabco blocks); only blocks whose catalog key matches size_keys are sized
  key     catalog key (config/families.json) the item prices and models as
  rank    which number wins when a model fits more than one (lower first; e.g. a new hand sink over an
          existing one) - the pick is reported so it can be overridden per store (item_overrides)

    python tools/master_items.py --refresh "<...>/PH MASTER EQUIPMENT LIST - W PRICING.xlsx"

re-reads the AutoQuotes export (existing items E# are written X#) and keeps the match / size / key /
rank / note already in the json.
"""
import argparse
import json
import os
import re

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
MASTER = os.path.join(ROOT, "config", "master_items.json")
KEEP = ("key", "match", "size", "size_keys", "rank", "note")


def _norm(s):
    return re.sub(r"[^A-Z0-9]", "", (s or "").upper())


def provided_by(item):
    """Item number prefix (Edgar): PH = provided by owner (Pizza Hut), X = existing (E in older drawings),
    plain number = AEQ supplies and installs. None when the item has no number."""
    s = str(item or "").strip().upper()
    if not s:
        return None
    if s.startswith("PH"):
        return "OWNER"
    return "EXISTING" if s[0] in "XE" else "AEQ"


def load_master(path=MASTER):
    if not os.path.isfile(path):
        return None
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def _item_no(v):
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    return "X" + s[1:] if s[:1].upper() == "E" else s


def refresh(xlsx, path=MASTER):
    from openpyxl import load_workbook
    ws = load_workbook(xlsx, data_only=True).worksheets[0]
    head = [str(c.value or "").strip() for c in ws[1]]
    col = {h: i for i, h in enumerate(head)}
    old = {it["item"]: it for it in (load_master(path) or {}).get("items", [])}
    items = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        n = r[col["ItemNo"]]
        if n in (None, ""):
            continue                                # continuation rows: warranty / option notes
        it = {"item": _item_no(n), "mfr": r[col["Mfr"]], "model": r[col["Model"]],
              "category": r[col["Category"]], "qty": r[col["Qty"]]}
        it["provided_by"] = provided_by(it["item"])
        it.update({k: v for k, v in old.get(it["item"], {}).items() if k in KEEP})
        items.append(it)
    data = {"_about": (load_master(path) or {}).get("_about", [__doc__.strip().splitlines()[0]]),
            "source": os.path.basename(xlsx), "items": items}
    with open(path, "w", encoding="utf-8", newline="\n") as f:
        f.write(json.dumps(data, indent=2, ensure_ascii=False) + "\n")
    return data


def _models(q):
    return [m for m in ((q.get("aq") or {}).get("model"), (q.get("kcl") or {}).get("model")) if m]


def _fits(it, q):
    """How q (an extracted block) is this master item: 'model', 'match' or 'size'; None if not."""
    norm_m = _norm(it.get("model"))
    for m in _models(q):
        n = _norm(m)
        if len(norm_m) >= 5 and len(n) >= 5 and (norm_m in n or n in norm_m):
            return "model"
    if len(norm_m) >= 5 and norm_m in _norm(q.get("block")):
        return "model"
    for rx in it.get("match", []):
        if any(re.search(rx, h, re.I) for h in _models(q) + [q.get("block") or ""]):
            return "match"
    if it.get("size") and q.get("bbox") and re.search(it.get("size_keys", "$^"), q.get("key") or ""):
        w, d = it["size"]
        bx, by = q["bbox"][2] - q["bbox"][0], q["bbox"][3] - q["bbox"][1]
        if (abs(bx - w) <= 3 and abs(by - d) <= 3) or (abs(bx - d) <= 3 and abs(by - w) <= 3):
            return "size"
    return None


def assign(out, store, master):
    """Number every extracted item from the master list (it replaces AutoQuotes' own numbers), set its
    catalog key and who provides it, and report what fits more than one number or none."""
    over = (store or {}).get("item_overrides", {})
    by_no = dict((it["item"], it) for it in master["items"])
    picks, missing = {}, {}
    for q in out["equipment"]:
        if q["status"] == "demo" or q["layer"].upper().startswith("FS-ELEC"):
            continue
        if q["handle"] in over:
            hits = [(by_no[over[q["handle"]]], "override")]
        else:
            hits = [(it, how) for it in master["items"] for how in [_fits(it, q)] if how]
            hits.sort(key=lambda h: (h[0].get("rank", 9), {"model": 0, "match": 1, "size": 2}[h[1]]))
        if not hits:                            # its AutoQuotes number is from the old numbering: drop it
            q["aq_item"], q["item"], q["provided_by"] = q.get("item"), None, None
            if q["key"] or _models(q):
                missing.setdefault(q.get("key") or _models(q)[0], []).append(q)
            continue
        it, how = hits[0]
        q["aq_item"], q["item"], q["master_match"] = q.get("item"), it["item"], how
        q["provided_by"] = provided_by(it["item"])
        if it.get("key"):
            q["key"] = it["key"]
        alts = sorted(set(h[0]["item"] for h in hits[1:]) - {it["item"]})
        if alts and how != "override":
            picks.setdefault((it["item"], tuple(alts)), []).append(q["handle"])
    for (n, alts), handles in sorted(picks.items()):
        out["warnings"].append("Item %s chosen for %d block(s) that could also be %s (handles %s): set "
                               "item_overrides in the store file if not." % (n, len(handles), "/".join(alts),
                                                                             ", ".join(handles)))
    # not on the master list = existing equipment at this store (Edgar): its own X number per model, after
    # the master's X numbers, priced like any existing item (moved, new utilities)
    nxt = 1 + max([int(m.group(1)) for it in master["items"]
                   for m in [re.match(r"X(\d+)$", it["item"])] if m] or [0])
    for k, qs in sorted(missing.items()):
        for q in qs:
            q["item"], q["provided_by"], q["master_match"] = "X%d" % nxt, "EXISTING", "existing, not on master"
        out["warnings"].append("Existing, not on the master list: X%d = %s (x%d)." % (nxt, k, len(qs)))
        nxt += 1


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--refresh", metavar="XLSX", required=True)
    a = ap.parse_args()
    d = refresh(a.refresh)
    print(MASTER, len(d["items"]), "items")
