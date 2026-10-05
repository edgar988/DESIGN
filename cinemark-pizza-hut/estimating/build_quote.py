#!/usr/bin/env python3
"""Customer construction quote PDF + INTERNAL estimate workbook for one store."""
import datetime as dt
import os

from reportlab.lib.units import inch
from reportlab.platypus import CondPageBreak, Image, KeepTogether, Paragraph, Spacer, Table, TableStyle

from . import quote_doc as Q

SCOPE_TEXT = {
    "GENERAL CONDITIONS": [
        "Mobilization for rough-in and for equipment set / finish; protection of adjacent theatre areas, "
        "dust barriers and walk-off mats.",
        "Permits and inspections (allowance), debris removal, final construction clean.",
    ],
    "DEMOLITION": [
        "Selective demolition within the Pizza Hut kitchen scope area per the layout.",
        "Disconnect and remove existing equipment shown for removal; cap and make safe existing utilities.",
    ],
    "CONSTRUCTION & FINISHES": [
        "Partitions, in-wall blocking for wall-hung sinks and shelving.",
        "Kitchen finishes: FRP wall panels, washable lay-in ceiling, sanitary flooring and cove base as noted.",
    ],
    "ELECTRICAL": [
        "Kitchen sub-panel and feeder; dedicated branch circuits, breakers and receptacles / direct connections "
        "for every piece of equipment per the electrical rough-in schedule.",
        "LED lighting in the kitchen area; circuits and disconnects for HVAC and exhaust.",
    ],
    "PLUMBING": [
        "Hot and cold water connections with shutoffs, direct and indirect waste, floor sinks with slab "
        "saw-cut and patch, point-of-use grease interceptor, mixing valve at hand sink, dish backflow "
        "preventer, per the plumbing rough-in schedule.",
    ],
    "HVAC": [
        "Dedicated kitchen cooling per the store heat-load assessment: inverter mini-split systems with "
        "concealed-duct indoor units, filtered returns, supply diffusers, line sets, condensate and "
        "wired controllers; general exhaust fan interlocked with the conveyor oven.",
    ],
    "EQUIPMENT SET & START-UP": [
        "Receive, inspect and stage owner-furnished Pizza Hut equipment at the store; set, level, connect "
        "and commission each piece (list below).",
    ],
    "TRAVEL & SUPERVISION": [
        "Project management, coordination, submittals and inspections; crew travel, lodging and per diem.",
    ],
}

EXCLUSIONS = [
    "Pizza Hut program equipment and smallwares: owner-furnished unless listed otherwise.",
    "PE-stamped engineering drawings (program add-on where the AHJ requires them).",
    "Hood, fire suppression and duct work: the package is ventless; any hood the AHJ requires is by change order.",
    "Asbestos or hazardous material abatement; concealed conditions; utility company charges.",
    "Electrical service or switchboard upgrades beyond the kitchen sub-panel and feeder shown.",
    "Low-voltage cabling (POS, network, music, cameras): rough-in boxes and conduit stubs only on request.",
    "Roof work beyond penetrations and flashing by the theatre's roofer of record (warranty).",
    "Work outside normal business hours unless stated; after-hours or weekend work is billed as an add.",
]


def section_text(sec, store, est):
    """The section's scope paragraph, from what is actually priced in it (and the store's construction
    narrative), so the words never promise work the numbers leave out. A store can still give its own
    text per section in scope.text."""
    scope = store.get("scope") or {}
    if (scope.get("text") or {}).get(sec):
        return [scope["text"][sec]]
    asm = set(l.get("assembly") for l in est["sections"][sec]["lines"])
    hv = store.get("hvac") or {}
    if sec == "DEMOLITION":
        out = []
        if "demo_lump_sum" in asm:
            out.append("Demolition in the existing room for the new kitchen, with debris removal and dumpsters.")
        if asm & {"demo_counter_front", "demo_counter_back"}:
            out.append("Remove the existing concession stand front counter and back counter / cabinets in the "
                       "footprint of the new layout; cap and make safe their utilities.")
        if asm & {"demo_equipment", "demo_partition", "demo_ceiling", "demo_flooring", "demo_casework"}:
            out.append("Selective demolition within the Pizza Hut kitchen scope area per the layout.")
        if asm & {"demo_equipment", "demo_make_safe"}:
            out.append("Disconnect and remove existing equipment shown for removal; cap and make safe existing "
                       "utilities.")
        return out
    if sec == "CONSTRUCTION & FINISHES":
        out = []
        if "con_cased_opening" in asm:
            out.append("New partition wall with a cased opening on either side for staff to pass in and out; "
                       "standard finish.")
        elif "con_partition" in asm:
            out.append("New partitions per the layout.")
        if "con_blocking" in asm:
            out.append("In-wall blocking for wall-hung sinks and shelving.")
        fin = [t for k, t in (("con_frp", "FRP wall panels"), ("con_act_washable", "washable lay-in ceiling"),
                              ("con_floor_quarry", "sanitary flooring"), ("con_floor_epoxy", "sanitary flooring"),
                              ("con_cove_base", "cove base")) if k in asm]
        if fin:
            fin = sorted(set(fin), key=fin.index)
            out.append("Kitchen finishes: %s." % (", ".join(fin[:-1]) + " and " + fin[-1] if len(fin) > 1 else fin[0]))
        return out
    if sec == "ELECTRICAL":
        sub = "el_subpanel" in asm
        out = [("Kitchen sub-panel and feeder; dedicated" if sub else "Dedicated") +
               " branch circuits%s, breakers and receptacles / direct connections for every piece of equipment "
               "per the electrical rough-in schedule." % ("" if sub else " from the existing panel")]
        if "el_lighting_sf" in asm:
            out.append("LED lighting in the kitchen area.")
        if "el_hvac_circuit" in asm:
            out.append("Circuit and disconnect for the HVAC condensing unit%s." % (
                " and exhaust fan" if hv.get("exhaust_fan") else ""))
        return out
    if sec == "HVAC" and "hv_split_system" in asm:
        heads = hv.get("heads") or 1
        out = ["Added cooling for the enclosed kitchen: %g-ton split system with %d indoor head%s, new refrigerant "
               "line sets, condensate drains and controls, start-up." % (hv.get("tons") or 0, heads,
                                                                       "s" if heads != 1 else "")]
        if "hv_crane" in asm:
            out.append("Condensing unit set by crane%s." % (
                " using the existing roof penetration" if hv.get("penetration_existing") else ""))
        if hv.get("exhaust_fan"):
            out.append("General exhaust fan interlocked with the conveyor oven.")
        return out
    if sec == "EQUIPMENT SET & START-UP" and asm & {"eq_receive_nashville", "eq_relocate"}:
        out = []
        if "eq_receive_nashville" in asm:
            out.append("Owner-furnished Pizza Hut equipment received, inspected and held at AEQ Nashville, then "
                       "delivered to site when the kitchen is ready.")
        if "eq_relocate" in asm:
            out.append("Existing equipment disconnected, protected, moved to its new location and reconnected "
                       "to new utilities.")
        out.append("Set, level, connect and commission each piece (list below).")
        return out
    return SCOPE_TEXT.get(sec, [])


EQUIPMENT_HEADS = [
    ("OWNER", "OWNER-FURNISHED EQUIPMENT (PH): RECEIVED, DELIVERED, SET &amp; CONNECTED", None),
    ("AEQ", "SUPPLIED BY AARON EQUIPMENT CO.: SET &amp; CONNECTED",
     "Equipment purchase is not included in this quote."),
    ("EXISTING", "EXISTING EQUIPMENT: RELOCATED &amp; RECONNECTED TO NEW UTILITIES", None),
]


def _qty(q, unit):
    if unit in ("LS", "allow"):
        return ""
    q = int(round(q)) if q >= 10 else round(q, 1)
    q = int(q) if float(q).is_integer() else q
    plural = "s" if q != 1 else ""
    unit = {"ea": "", "LF": " LF", "SF": " SF", "pull": " pull" + plural, "trip": " trip" + plural,
            "hr": " hr"}.get(unit, " " + unit)
    return " (%s%s)" % (format(q, ","), unit)


def scope_detail(lines):
    """Customer-safe scope itemization: descriptions + quantities, never prices or hours."""
    out = []
    for l in lines:
        d = l["desc"].replace("&", "&amp;")
        if out and d[1:2].islower():            # not acronyms: LED, FRP, HVAC
            d = d[0].lower() + d[1:]
        out.append(d + _qty(l["qty"], l["unit"]))
    return "; ".join(out)


def _fmt_date(d):
    return d.strftime("%B %d, %Y").replace(" 0", " ")


def quote_numbers(store, rev="R0", today=None):
    today = today or dt.date.today()
    no = "SSG-2026-CNK-PH01-%03d-C1" % store["theatre_no"]
    slug = store.get("slug") or "%03d" % store["theatre_no"]
    fname = "%s_%s_CINEMARK_%s_PH-CONSTRUCTION_%s_%s_%s" % (
        today.strftime("%y"), store["state"], slug, no, rev, today.strftime("%m-%d-%y"))
    return no, fname


def build_pdf(path, store, program, takeoff, est, sched, gantt_png, rev="R0", today=None):
    today = today or dt.date.today()
    cust, con, prog = program["customer"], program["contractor"], program["program"]
    no, _ = quote_numbers(store, rev, today)
    loc = "#%d %s" % (store["theatre_no"], store["theatre_name"].split(" (")[0])
    doc = Q.make_doc(path, no, "%s · %s" % (prog["name"], loc), con, rev)
    s = []
    s.append(Paragraph("STORE CONSTRUCTION QUOTE · %s · PROGRAM %s" % (
        prog["tier"].split(":")[0].upper(), prog["proposal_ref"]), Q.eyebrow))
    s.append(Paragraph("Pizza Hut Kitchen Conversion · Cinemark %s" % loc, Q.h1))
    s.append(Paragraph(store["address"], Q.h1sub))
    s.append(Spacer(1, 6))
    s.append(Q.red_rule())
    s.append(Spacer(1, 4))
    heat = takeoff["heat_load"]
    has_hvac = "HVAC" in est["sections"]
    s.append(Q.panels(
        "CUSTOMER",
        [("Customer:", cust["name"]), ("Address:", cust["address"]), ("Attn:", cust["attn"]),
         ("cc:", cust["cc"]), ("Quote No.:", "%s %s" % (no, rev)), ("Date:", _fmt_date(today)),
         ("Valid For:", "%d days from Quote Date" % prog["validity_days"])],
        "JOB SITE",
        [("Theatre:", "Cinemark %s" % loc), ("Address:", store["address"]),
         ("Scope area:", "%s SF Pizza Hut kitchen" % format(round(takeoff["quantities"]["room_sf"]), ",")),
         ("Basis:", "Store survey, %s%s" % (
             ("layout " + store.get("layout_dxf", "")) if takeoff.get("layout_used")
             else ", ".join(d for d in store.get("existing_docs", ["equipment package list"])
                            if "Heat-Load" not in d),
             ", heat-load assessment SSG-2026-CNK-PH01-%03d-M1" % store["theatre_no"] if has_hvac else "")),
         ("Connected load:", "%.1f kW equipment%s" % (
             heat["connected_kw"], "; %.1f tons added cooling" % heat["minisplit_tons"] if has_hvac else ""))]))
    s.append(Paragraph("CONTRACTOR", Q.sechead))
    s.append(Q.kv([("Contractor:", con["entity"]), ("Address:", con["address"]), ("Contact:", con["contact"]),
                   ("Phone / Email:", "%s · %s" % (con["phone"], con["email"])),
                   ("Licenses:", con["licenses"])]))

    s.append(Paragraph("PRICING", Q.sechead))
    rows = []
    for sec in est["section_order"]:
        detail = scope_detail(est["sections"][sec]["lines"]) if sec != "TRAVEL & SUPERVISION" else ""
        rows.append(("<b>%s</b><br/>%s%s" % (
            sec.title().replace("&", "&amp;").replace("Hvac", "HVAC"),
            " ".join(section_text(sec, store, est)),
            ("<br/><font color='#666666'>Includes: %s.</font>" % detail) if detail else ""),
            est["customer_sections"][sec]))
    s.append(Q.price_table(rows, "TOTAL · LUMP SUM, ALL TRAVEL INCLUDED", est["total"]))
    s.append(Paragraph("Section amounts include receiving, coordination, supervision and travel as allocated. "
                       "Survey, design and rough-in drawings are billed under the program proposal (Tier 2) "
                       "and are not included here.", Q.sub))

    def _eq(rows):
        n = {}                                  # one line per item: blocks of the same item add up
        for r in rows:
            k = (r.get("item"), r["mfr"], r["model"], r["description"])
            n[k] = n.get(k, 0) + r["qty"]
        return [("%s(%d) %s %s, %s" % (("Item %s: " % k[0]) if k[0] else "", q, k[1], k[2], k[3])).replace("&", "&amp;")
                for k, q in n.items()]

    s.append(CondPageBreak(2.5 * inch))
    if any(r.get("provided_by") for r in takeoff["schedule"]):     # numbered from the PH master list
        for src, head, note in EQUIPMENT_HEADS:
            rows = [r for r in takeoff["schedule"] if (r.get("provided_by") or "AEQ") == src]
            if rows:
                s.append(Paragraph(head, Q.sechead))
                s.extend(Q.bullets(_eq(rows)))
                if note:
                    s.append(Paragraph(note, Q.sub))
    else:
        s.append(Paragraph("OWNER-FURNISHED EQUIPMENT SET &amp; CONNECTED", Q.sechead))
        s.extend(Q.bullets(_eq(takeoff["schedule"])))

    s.append(CondPageBreak(3.2 * inch))
    s.append(Paragraph("SCHEDULE", Q.sechead))
    s.append(Paragraph("Notice to proceed %s · on site %s · turnover %s · about %d working days on "
                       "site, %.1f calendar weeks overall. %s Dates move with permit issue and owner equipment "
                       "delivery." % (_fmt_date(sched["ntp"]), _fmt_date(sched["site_start"]),
                                      _fmt_date(sched["turnover"]), sched["site_working_days"],
                                      sched["calendar_weeks"], sched["note"]), Q.base))
    if gantt_png and os.path.exists(gantt_png):
        from PIL import Image as PILImage
        iw, ih = PILImage.open(gantt_png).size
        w = Q.CW
        s.append(Image(gantt_png, width=w, height=w * ih / iw))

    s.append(KeepTogether([Paragraph("INCLUSIONS", Q.sechead)] + Q.bullets([
        "All labor, materials, travel, lodging and per diem for the scope above; work self-performed by "
        "Spirit Services Group crews.",
        "Rough-in per the store's electrical and plumbing rough-in schedules; final connections within %d ft "
        "of each equipment connection point." % program["rough_in_defaults"]["point_of_connection_ft"],
        "Start-up and commissioning of installed systems; Hobart dish machine start-up by the local Hobart "
        "service office for warranty.",
        "One-year workmanship warranty on SSG-performed installation from date of substantial completion.",
    ])))
    excl = list(EXCLUSIONS)
    if "el_subpanel" not in set(l.get("assembly") for l in est["lines"]):
        excl[excl.index(EXCLUSIONS[4])] = ("Electrical service, panel or switchboard upgrades: circuits are fed from "
                                           "the existing panel, assumed to have spare capacity and breaker spaces.")
    if not has_hvac:
        excl[excl.index(EXCLUSIONS[6])] = "HVAC: none included; the area stays on the theatre's existing system."
    s.append(KeepTogether([Paragraph("EXCLUSIONS / ASSUMPTIONS", Q.sechead)] + Q.bullets(
        excl + ["Work performed during normal business hours with the kitchen area vacated; theatre "
                      "remains open." if store.get("schedule", {}).get("work_window", "day") == "day"
                      else "After-hours work included as noted in the schedule.",
                      "Sales / use tax per applicable jurisdiction; non-union pricing; quoted materials subject "
                      "to manufacturer price increases.",
                      "Alterations or deviations from scope causing extra cost proceed only with a signed "
                      "change order."]
        + ["Verify at site: %s." % v for v in store.get("customer_verify", [])])))
    s.append(Spacer(1, 8))
    s.append(Q.acceptance(cust["name"], con["entity"], con["contact"].split(",")[0],
                          "Scheduling begins upon receipt of signed acceptance or written confirmation to %s."
                          % con["email"]))
    doc.build(s)
    return path


def build_internal_xlsx(path, store, takeoff, est, sched):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    hdr = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="222222")
    red = Font(bold=True, color="C1121F")

    def sheet(ws, headers, rows, widths=None):
        ws.append(headers)
        for c in ws[1]:
            c.font, c.fill = hdr, fill
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths or [], 1):
            ws.column_dimensions[chr(64 + i)].width = w
        ws.freeze_panes = "A2"

    ws = wb.active
    ws.title = "Summary"
    ws.append(["INTERNAL - NOT FOR CUSTOMER", store["store_id"], store["theatre_name"]])
    ws["A1"].font = red
    ws.append([])
    ws.append(["Section", "Hours", "True cost", "Build sell", "Customer sell", "Margin %"])
    for c in ws[3]:
        c.font, c.fill = hdr, fill
    for sec in est["section_order"]:
        d = est["sections"][sec]
        cs = est["customer_sections"][sec]
        ws.append([sec, round(d["hrs"], 1), round(d["cost"], 2), round(d["sell"], 2), cs,
                   round((cs - d["cost"]) / cs * 100, 1) if cs else 0])
    ws.append([])
    for k, v in [("Build subtotal", est["subtotal_sell"]), ("Contingency (in sections)", est["contingency"]),
                 ("TOTAL (customer)", est["total"]), ("True cost", est["true_cost"]),
                 ("Margin %", est["margin_pct"]), ("Margin check", est["markup_check"]),
                 ("Field hours", est["labor"]["field_hrs"]), ("Crew", est["labor"]["crew"]),
                 ("Site days (labor basis)", est["labor"]["site_days"]),
                 ("Travel days", est["labor"]["travel_days"]), ("Lodging nights", est["labor"]["nights"]),
                 ("Billed rate", est["rates_used"]["labor_billed_hr"]),
                 ("True labor cost/hr", est["rates_used"]["labor_true_cost_hr"])]:
        ws.append([k, v])
    ws.column_dimensions["A"].width = 34
    for col in "BCDEF":
        ws.column_dimensions[col].width = 15

    sheet(wb.create_sheet("Lines"),
          ["Section", "Assembly", "Description", "Qty", "Unit", "Hours", "Material", "Sub", "Equipment",
           "True cost", "Sell", "Basis", "Note"],
          [[l["section"], l.get("assembly", ""), l["desc"], l["qty"], l["unit"], l["hrs"], l["mat"], l["sub"],
            l["eqp"], l["cost"], l["sell"], l["basis"], l.get("note", "")] for l in est["lines"]],
          [24, 22, 60, 9, 7, 9, 11, 10, 11, 12, 12, 40, 34])
    sheet(wb.create_sheet("Rough-In Electrical"),
          ["Ckt", "Item", "Key", "Volts", "Phase", "Amps", "Breaker A", "Poles", "Conn", "NEMA", "Assembly",
           "Home run LF", "Height AFF", "Basis"],
          [[c["circuit"], c["item"], c["key"], c["volts"], c["phase"], c["amps"], c["breaker_a"], c["poles"],
            c["conn"], c["nema"], c["assembly"], c["home_run_lf"], c["height_aff"], c["breaker_basis"]]
           for c in takeoff["circuits"]], [6, 6, 20, 7, 7, 7, 10, 7, 8, 10, 18, 12, 10, 36])
    sheet(wb.create_sheet("Equipment"),
          ["Item", "Qty", "Mfr", "Model", "Description", "Electrical", "Plumbing", "Sens Btu/h", "Lat Btu/h",
           "Verified", "Flags"],
          [[r.get("item"), r["qty"], r["mfr"], r["model"], r["description"], r.get("elec", ""),
            r.get("plumb", ""), r["heat_sensible"], r["heat_latent"], r["verified"], " | ".join(r["flags"])]
           for r in takeoff["schedule"]], [6, 5, 18, 16, 40, 34, 28, 11, 10, 9, 70])
    h = takeoff["heat_load"]
    sheet(wb.create_sheet("Heat Load"), ["Quantity", "Value"], [[k, v] for k, v in h.items()], [28, 16])
    sheet(wb.create_sheet("Schedule"), ["Task", "Start", "End", "Days", "Kind"],
          [[t["name"], t["start"], t["end"], t["days"], t["kind"]] for t in sched["tasks"]], [62, 12, 12, 7, 8])
    flags = [["Unverified catalog entry", k] for k in takeoff["unverified"]]
    flags += [["Takeoff warning", w] for w in takeoff["warnings"]]
    for r in takeoff["schedule"]:
        flags += [["Equipment flag " + r["key"], f] for f in r["flags"]]
    flags += [["Store note", n] for n in store.get("notes", [])]
    sheet(wb.create_sheet("Flags"), ["Type", "Detail"], flags, [34, 120])
    for w in wb.worksheets:
        for row in w.iter_rows():
            for c in row:
                c.alignment = Alignment(vertical="top", wrap_text=False)
    wb.save(path)
    return path
